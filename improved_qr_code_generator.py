#!/usr/bin/env python3
"""
Improved Link Extractor and QR Code Generator

This script extracts links from various text-based files, stores them in a CSV,
generates QR codes, and creates a PDF with 4 codes per page.
"""

import os
import re
import csv
import yaml
import qrcode
import argparse
import datetime
from urllib.parse import urlparse
import uuid
from typing import List, Dict, Optional, Set, Any, Tuple
from pathlib import Path

# Import optional modules with error handling
try:
    from fpdf import FPDF
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    import PyPDF2
    PYPDF2_AVAILABLE = True
except ImportError:
    PYPDF2_AVAILABLE = False

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import pandas as pd
    PANDAS_AVAILABLE = True
except ImportError:
    PANDAS_AVAILABLE = False

try:
    from bs4 import BeautifulSoup
    BEAUTIFULSOUP_AVAILABLE = True
except ImportError:
    BEAUTIFULSOUP_AVAILABLE = False


class LinkExtractor:
    """Extract links from various file types."""
    
    def __init__(self, config: Dict[str, Any]):
        """
        Initialize the LinkExtractor with configuration.
        
        Args:
            config: Configuration dictionary
        """
        self.config = config
        self.url_pattern = re.compile(r'http[s]?://(?:[a-zA-Z]|[0-9]|[$-_@.&+]|[!*\\(\\),]|(?:%[0-9a-fA-F][0-9a-fA-F]))+')
    
    def extract_from_file(self, file_path: str) -> Set[str]:
        """
        Extract links from a file based on its extension.
        
        Args:
            file_path: Path to the file
            
        Returns:
            Set of unique URLs found in the file
        """
        file_extension = os.path.splitext(file_path)[1].lower()
        
        extractors = {
            '.txt': self._extract_from_text,
            '.md': self._extract_from_text,
            '.csv': self._extract_from_csv,
            '.pdf': self._extract_from_pdf,
            '.docx': self._extract_from_docx,
            '.doc': self._extract_from_text,  # Fallback for .doc files
            '.html': self._extract_from_html,
            '.htm': self._extract_from_html,
            '.yaml': self._extract_from_yaml,
            '.yml': self._extract_from_yaml,
            '.json': self._extract_from_json
        }
        
        extractor = extractors.get(file_extension, self._extract_from_text)
        try:
            return extractor(file_path)
        except Exception as e:
            print(f"Error extracting links from {file_path}: {e}")
            return set()
    
    def _extract_from_text(self, file_path: str) -> Set[str]:
        """Extract links from text-based files."""
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
                content = file.read()
            return self._find_urls(content)
        except Exception as e:
            print(f"Error reading text file {file_path}: {e}")
            return set()
    
    def _extract_from_csv(self, file_path: str) -> Set[str]:
        """Extract links from CSV files."""
        if not PANDAS_AVAILABLE:
            print("Pandas is not installed. Falling back to text extraction for CSV.")
            return self._extract_from_text(file_path)
        
        try:
            df = pd.read_csv(file_path)
            urls = set()
            for column in df.columns:
                for value in df[column].astype(str):
                    urls.update(self._find_urls(value))
            return urls
        except Exception as e:
            print(f"Error processing CSV file {file_path}: {e}")
            return self._extract_from_text(file_path)
    
    def _extract_from_pdf(self, file_path: str) -> Set[str]:
        """Extract links from PDF files."""
        if not PYPDF2_AVAILABLE:
            print("PyPDF2 is not installed. Falling back to text extraction for PDF.")
            return self._extract_from_text(file_path)
        
        try:
            pdf_file = open(file_path, 'rb')
            pdf_reader = PyPDF2.PdfReader(pdf_file)
            
            urls = set()
            for page in pdf_reader.pages:
                text = page.extract_text()
                if text:
                    urls.update(self._find_urls(text))
            
            pdf_file.close()
            return urls
        except Exception as e:
            print(f"Error processing PDF file {file_path}: {e}")
            return set()
    
    def _extract_from_docx(self, file_path: str) -> Set[str]:
        """Extract links from DOCX files."""
        if not DOCX_AVAILABLE:
            print("python-docx is not installed. Falling back to text extraction for DOCX.")
            return self._extract_from_text(file_path)
        
        try:
            doc = docx.Document(file_path)
            content = "\n".join([para.text for para in doc.paragraphs])
            return self._find_urls(content)
        except Exception as e:
            print(f"Error processing DOCX file {file_path}: {e}")
            return set()
    
    def _extract_from_html(self, file_path: str) -> Set[str]:
        """Extract links from HTML files."""
        if not BEAUTIFULSOUP_AVAILABLE:
            print("BeautifulSoup is not installed. Falling back to text extraction for HTML.")
            return self._extract_from_text(file_path)
        
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
                content = file.read()
            
            soup = BeautifulSoup(content, 'html.parser')
            
            # Extract URLs from href attributes
            urls = set()
            for a_tag in soup.find_all('a', href=True):
                href = a_tag['href']
                if href.startswith('http'):
                    urls.add(href)
            
            # Also find URLs in the text using regex for any missed links
            urls.update(self._find_urls(content))
            
            return urls
        except Exception as e:
            print(f"Error processing HTML file {file_path}: {e}")
            return self._extract_from_text(file_path)
    
    def _extract_from_yaml(self, file_path: str) -> Set[str]:
        """Extract links from YAML files."""
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
                content = file.read()
                
            # First, try to find URLs in the raw text
            urls = self._find_urls(content)
            
            # Then try to parse YAML and search for URLs in values
            try:
                yaml_data = yaml.safe_load(content)
                if yaml_data:
                    for key, value in yaml_data.items():
                        if isinstance(value, str):
                            urls.update(self._find_urls(value))
            except Exception as yaml_e:
                print(f"Error parsing YAML structure in {file_path}: {yaml_e}")
                # Already found URLs from raw text, so continue
                
            return urls
        except Exception as e:
            print(f"Error processing YAML file {file_path}: {e}")
            return set()
    
    def _extract_from_json(self, file_path: str) -> Set[str]:
        """Extract links from JSON files."""
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
                content = file.read()
                
            # Just find URLs in the raw text
            urls = self._find_urls(content)
            return urls
        except Exception as e:
            print(f"Error processing JSON file {file_path}: {e}")
            return set()
    
    def _find_urls(self, text: str) -> Set[str]:
        """Find all URLs in a text string."""
        return set(self.url_pattern.findall(text))


class CSVProcessor:
    """Process and manage URLs in CSV format with timestamps."""
    
    def __init__(self, output_path: str):
        """
        Initialize the CSV processor.
        
        Args:
            output_path: Directory to save the CSV file
        """
        self.output_path = output_path
        self.csv_path = os.path.join(output_path, "urls.csv")
        self.urls = set()  # For tracking unique URLs
        
    def add_urls(self, urls: Set[str]) -> List[str]:
        """
        Add unique URLs to the CSV.
        
        Args:
            urls: Set of URLs to add
            
        Returns:
            List of newly added URLs
        """
        # Create directory if it doesn't exist
        os.makedirs(self.output_path, exist_ok=True)
        
        # Read existing URLs if file exists
        existing_urls = set()
        if os.path.exists(self.csv_path):
            with open(self.csv_path, 'r', encoding='utf-8') as f:
                reader = csv.DictReader(f)
                for row in reader:
                    existing_urls.add(row['url'])
        
        # Combine with already tracked URLs
        self.urls.update(existing_urls)
        
        # Find new URLs
        new_urls = [url for url in urls if url not in self.urls]
        timestamp = datetime.datetime.now().isoformat()
        
        # Add new URLs to CSV
        file_exists = os.path.exists(self.csv_path)
        with open(self.csv_path, 'a', encoding='utf-8', newline='') as f:
            fieldnames = ['url', 'timestamp', 'version', 'binary_data']
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            
            if not file_exists:
                writer.writeheader()
            
            for url in new_urls:
                writer.writerow({
                    'url': url,
                    'timestamp': timestamp,
                    'version': '',  # Will be filled later
                    'binary_data': ''  # Will be filled later
                })
        
        # Add new URLs to the tracking set
        self.urls.update(new_urls)
        
        return new_urls
    
    def update_qr_data(self, url: str, version: int, binary_data: str) -> None:
        """
        Update QR code data for a URL in the CSV.
        
        Args:
            url: URL to update
            version: QR code version
            binary_data: Binary representation of QR code
        """
        if not os.path.exists(self.csv_path):
            return
        
        # Read all rows
        rows = []
        with open(self.csv_path, 'r', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            for row in reader:
                if row['url'] == url:
                    row['version'] = str(version)
                    row['binary_data'] = binary_data
                rows.append(row)
        
        # Write back all rows
        with open(self.csv_path, 'w', encoding='utf-8', newline='') as f:
            fieldnames = ['url', 'timestamp', 'version', 'binary_data']
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(rows)
    
    def get_urls(self) -> List[Dict[str, str]]:
        """
        Get all URLs and their data from the CSV.
        
        Returns:
            List of dictionaries with URL data
        """
        if not os.path.exists(self.csv_path):
            return []
        
        urls_data = []
        with open(self.csv_path, 'r', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            for row in reader:
                urls_data.append(dict(row))
        
        return urls_data


class QRCodeGenerator:
    """Generate QR codes from URLs."""
    
    def __init__(self, config: Dict[str, Any], url_manager=None):
        """
        Initialize the QRCodeGenerator with configuration.
        
        Args:
            config: Configuration dictionary
            url_manager: Optional URL manager instance
        """
        self.config = config
        self.error_correction = getattr(qrcode.constants, 
                                      config.get('qr_error_correction', 'ERROR_CORRECT_L'))
        self.box_size = config.get('qr_box_size', 10)
        self.border = config.get('qr_border', 4)
        self.url_manager = url_manager
    
    def generate_qr_code(self, url: str, output_path: str, filename: str = "qrcode") -> Dict[str, str]:
        """
        Generate QR code for a URL and save as image and binary file.
        
        Args:
            url: URL to encode in QR code
            output_path: Directory to save the QR code files
            filename: Base filename for output files
            
        Returns:
            Dictionary containing paths to generated files
        """
        os.makedirs(output_path, exist_ok=True)
        
        # Create QR code
        qr = qrcode.QRCode(
            version=None,  # Auto-determine
            error_correction=self.error_correction,
            box_size=self.box_size,
            border=self.border,
        )
        qr.add_data(url)
        qr.make(fit=True)
        
        img = qr.make_image(fill_color="black", back_color="white")
        
        # Save image
        image_path = os.path.join(output_path, f"{filename}.png")
        img.save(image_path)
        
        # Get and save binary data
        binary_data = self._get_binary_data(qr)
        binary_path = os.path.join(output_path, f"{filename}.bin")
        self._save_binary_data(qr, url, binary_path)
        
        # Save URL to markdown file
        url_path = os.path.join(output_path, "url.md")
        with open(url_path, 'w', encoding='utf-8') as f:
            f.write(f"# QR Code Link\n\n{url}\n")
        
        # Update URL manager if available
        if self.url_manager:
            self.url_manager.mark_as_processed(url, qr.version, binary_data)
        
        return {
            'image': image_path,
            'binary': binary_path,
            'url': url_path
        }
    
    def _get_binary_data(self, qr: qrcode.QRCode) -> str:
        """Extract binary data from QR code."""
        data = qr.get_matrix()
        binary_data = ""
        for row in data:
            for cell in row:
                binary_data += "1" if cell else "0"
        return binary_data

class PDFGenerator:
    """Generate PDF with QR codes, 4 codes per page."""
    
    def __init__(self, config: Dict[str, Any]):
        """
        Initialize the PDF generator with configuration.
        
        Args:
            config: Configuration dictionary
        """
        self.config = config
    
    def create_pdf(self, qr_data: List[Dict[str, Any]], output_path: str) -> Optional[str]:
        """
        Create PDF with QR codes, 4 per page.
        
        Args:
            qr_data: List of dictionaries containing QR code data
            output_path: Path to save the PDF file
            
        Returns:
            Path to the created PDF file or None if failed
        """
        if not PDF_AVAILABLE:
            print("fpdf is not installed. PDF creation skipped.")
            return None
        
        if not qr_data:
            print("No QR codes to include in PDF.")
            return None
        
        try:
            pdf = FPDF()
            pdf.set_auto_page_break(auto=True, margin=15)
            
            # Process 4 QR codes per page
            for i in range(0, len(qr_data), 4):
                pdf.add_page()
                
                # Get batch of up to 4 codes
                batch = qr_data[i:i+4]
                
                for j, data in enumerate(batch):
                    # Calculate position (2x2 grid)
                    row = j // 2
                    col = j % 2
                    
                    x = 30 + col * 85  # Horizontal spacing
                    y = 30 + row * 120  # Vertical spacing
                    
                    # Add URL as title
                    pdf.set_font("Arial", "B", 10)
                    pdf.set_xy(x, y)
                    url = data['url']
                    # Truncate URL if too long
                    if len(url) > 40:
                        display_url = url[:37] + "..."
                    else:
                        display_url = url
                    pdf.cell(70, 10, display_url, ln=True, align='C')
                    
                    # Add QR code image
                    pdf.image(data['image_path'], x + 15, y + 15, w=40)
                    
                    # Add timestamp
                    pdf.set_font("Arial", "", 8)
                    pdf.set_xy(x, y + 60)
                    timestamp = data.get('timestamp', '')
                    try:
                        # Format timestamp nicely if possible
                        dt = datetime.datetime.fromisoformat(timestamp)
                        formatted_time = dt.strftime("%Y-%m-%d %H:%M:%S")
                    except:
                        formatted_time = timestamp
                    pdf.cell(70, 5, f"Generated: {formatted_time}", ln=True, align='C')
                    
                    # Add version
                    pdf.set_xy(x, y + 65)
                    version = data.get('version', '')
                    pdf.cell(70, 5, f"Version: {version}", ln=True, align='C')
            
            pdf_path = os.path.join(output_path, "qr_codes.pdf")
            pdf.output(pdf_path)
            return pdf_path
        
        except Exception as e:
            print(f"Error creating PDF: {e}")
            return None

from url_manager import URLManager
class LinkProcessorApp:
    """Main application class for processing links and generating outputs."""
    
    def __init__(self, config_path: Optional[str] = None, config: Optional[Dict[str, Any]] = None):
        """
        Initialize the application with optional configuration file or dictionary.
        
        Args:
            config_path: Path to configuration YAML file
            config: Configuration dictionary (overrides config_path if provided)
        """
        if config:
            self.config = config
        else:
            self.config = self._load_config(config_path)
        
        self.link_extractor = LinkExtractor(self.config)
    
    def _load_config(self, config_path: Optional[str]) -> Dict[str, Any]:
        """Load configuration from YAML file or use defaults."""
        default_config = {
            'output_pdf': True,
            'output_raw': True,
            'qr_error_correction': 'ERROR_CORRECT_L',
            'qr_box_size': 10,
            'qr_border': 4,
            'generate_descriptive_titles': True
        }
        
        if config_path and os.path.exists(config_path):
            try:
                with open(config_path, 'r', encoding='utf-8') as f:
                    user_config = yaml.safe_load(f)
                if user_config:
                    default_config.update(user_config)
            except Exception as e:
                print(f"Error loading config from {config_path}: {e}")
                print("Using default configuration.")
        
        return default_config
    
    def process_file(self, input_file: str, output_dir: str = "output", url_manager=None) -> Dict[str, Any]:
        """
        Process input file to extract links and generate QR codes.
        
        Args:
            input_file: Path to input file
            output_dir: Base directory for outputs
            url_manager: Optional URL manager instance
            
        Returns:
            Dictionary with processing results
        """
        print(f"Processing file: {input_file}")
        
        # Helper function to generate title from URL
        def generate_title(url: str) -> str:
            """
            Generate a descriptive title for a URL.
            
            Args:
                url: URL to generate title for
                
            Returns:
                Generated title
            """
            if not self.config.get('generate_descriptive_titles', True):
                return url
            
            # Extract domain and path components
            from urllib.parse import urlparse
            parsed_url = urlparse(url)
            domain = parsed_url.netloc.replace('www.', '')
            
            path_parts = [p for p in parsed_url.path.split('/') if p]
            
            # Format title
            if not path_parts:
                return domain
            
            # Use last meaningful path component as title
            import re
            slug = path_parts[-1].replace('-', ' ').replace('_', ' ')
            
            # Clean up the slug
            slug = re.sub(r'\.\w+$', '', slug)  # Remove file extensions
            
            if not slug:
                return domain
            
            return f"{slug.title()} - {domain}"
        
        # Create or use URL manager
        if not url_manager:
            # We need to define URLManager class or import it
            class URLManager:
                """Manage URLs and their associated metadata to prevent duplicates"""
                
                def __init__(self):
                    """Initialize the URL manager with empty collections"""
                    self.urls = {}  # Dictionary mapping URLs to their metadata
                    self.url_hashes = set()  # Set of URL hashes for quick lookup
                    from datetime import datetime
                    self.datetime = datetime
                    import uuid
                    self.uuid = uuid
                
                def add_url(self, url: str) -> Dict:
                    """
                    Add a URL to the manager if it doesn't exist.
                    
                    Args:
                        url: The URL to add
                        
                    Returns:
                        Dictionary with URL metadata
                    """
                    # Create a unique hash for the URL
                    url_hash = self._hash_url(url)
                    
                    # Check if URL already exists
                    if url_hash in self.url_hashes:
                        return self.urls[url]
                    
                    # Add new URL
                    url_id = str(self.uuid.uuid4())
                    from urllib.parse import urlparse
                    parsed_url = urlparse(url)
                    
                    metadata = {
                        'url': url,
                        'id': url_id,
                        'domain': parsed_url.netloc,
                        'path': parsed_url.path,
                        'query': parsed_url.query,
                        'timestamp': None,  # Will be set when QR code is generated
                        'version': None,    # Will be set when QR code is generated
                        'qr_data': None,    # Will be set when QR code is generated
                        'processed': False  # Flag to indicate if QR code has been generated
                    }
                    
                    self.urls[url] = metadata
                    self.url_hashes.add(url_hash)
                    
                    return metadata
                
                def _hash_url(self, url: str) -> str:
                    """
                    Create a hash of the URL for duplicate detection.
                    
                    Args:
                        url: The URL to hash
                        
                    Returns:
                        A string hash of the URL
                    """
                    # This is a simple hash, but we could use a more robust method
                    return url
                
                def get_unprocessed_urls(self) -> List[Dict]:
                    """
                    Get all URLs that haven't been processed yet.
                    
                    Returns:
                        List of URL metadata dictionaries
                    """
                    return [metadata for url, metadata in self.urls.items() if not metadata['processed']]
                
                def mark_as_processed(self, url: str, version: int, qr_data: str) -> None:
                    """
                    Mark a URL as processed with QR code data.
                    
                    Args:
                        url: The URL that was processed
                        version: QR code version number
                        qr_data: Binary QR code data
                    """
                    if url in self.urls:
                        self.urls[url]['processed'] = True
                        self.urls[url]['timestamp'] = self.datetime.now().isoformat()
                        self.urls[url]['version'] = version
                        self.urls[url]['qr_data'] = qr_data
                
                def export_to_csv(self, output_path: str) -> str:
                    """
                    Export all URLs and their metadata to a CSV file.
                    
                    Args:
                        output_path: Directory to save the CSV file
                        
                    Returns:
                        Path to the created CSV file
                    """
                    import os
                    import csv
                    csv_path = os.path.join(output_path, "urls.csv")
                    
                    with open(csv_path, 'w', newline='', encoding='utf-8') as f:
                        writer = csv.writer(f)
                        writer.writerow(['url', 'timestamp', 'version', 'binary_data'])
                        
                        for url, metadata in self.urls.items():
                            writer.writerow([
                                url,
                                metadata.get('timestamp', ''),
                                metadata.get('version', ''),
                                metadata.get('qr_data', '')
                            ])
                    
                    return csv_path
                    
            url_manager = URLManager()
        
        # Extract base filename without extension
        base_filename = os.path.splitext(os.path.basename(input_file))[0]
        
        # Create output directory
        file_output_dir = os.path.join(output_dir, base_filename)
        os.makedirs(file_output_dir, exist_ok=True)
        
        # Extract links
        links = self.link_extractor.extract_from_file(input_file)
        print(f"Found {len(links)} unique links")
        
        # Add URLs to manager
        for url in links:
            url_manager.add_url(url)
        
        # Get unprocessed URLs
        unprocessed_urls = url_manager.get_unprocessed_urls()
        print(f"Processing {len(unprocessed_urls)} new links")
        
        # Define a modified QR code generator that works with our URL manager
        class EnhancedQRGenerator(QRCodeGenerator):
            def __init__(self, config, url_manager):
                super().__init__(config)
                self.url_manager = url_manager
                
            def generate_qr_code(self, url, output_path, filename="qrcode"):
                os.makedirs(output_path, exist_ok=True)
                
                # Create QR code
                qr = qrcode.QRCode(
                    version=None,  # Auto-determine
                    error_correction=self.error_correction,
                    box_size=self.box_size,
                    border=self.border,
                )
                qr.add_data(url)
                qr.make(fit=True)
                
                img = qr.make_image(fill_color="black", back_color="white")
                
                # Save image
                image_path = os.path.join(output_path, f"{filename}.png")
                img.save(image_path)
                
                # Get binary data
                data = qr.get_matrix()
                binary_data = ""
                for row in data:
                    for cell in row:
                        binary_data += "1" if cell else "0"
                
                # Save binary data
                binary_path = os.path.join(output_path, f"{filename}.bin")
                additional_info = f"Link: {url}\nVersion: {qr.version}\n"
                data_file_content = additional_info.encode('utf-8') + binary_data.encode('utf-8')
                
                with open(binary_path, "wb") as f:
                    f.write(data_file_content)
                
                # Save URL to markdown file
                url_path = os.path.join(output_path, "url.md")
                with open(url_path, 'w', encoding='utf-8') as f:
                    f.write(f"# QR Code Link\n\n{url}\n")
                
                # Update URL manager
                self.url_manager.mark_as_processed(url, qr.version, binary_data)
                
                return {
                    'image': image_path,
                    'binary': binary_path,
                    'url': url_path
                }
        
        # Create QR generator with URL manager
        qr_generator = EnhancedQRGenerator(self.config, url_manager)
        
        # Process each unprocessed link
        qr_data = []
        for metadata in unprocessed_urls:
            url = metadata['url']
            
            # Use URL ID for directory name to ensure uniqueness
            dir_name = f"url_{metadata['id']}"
            url_dir = os.path.join(file_output_dir, dir_name)
            
            # Generate QR code
            files = qr_generator.generate_qr_code(url, url_dir)
            
            # Generate title using our local function
            title = generate_title(url)
            
            qr_data.append({
                'url': url,
                'title': title,
                'directory': url_dir,
                'files': files,
                'metadata': metadata
            })
        
        # Export URL data to CSV
        csv_path = url_manager.export_to_csv(file_output_dir)
        print(f"URL data exported to: {csv_path}")
        
        # Generate outputs if configured and links exist
        results = {
            'input_file': input_file,
            'output_dir': file_output_dir,
            'links_found': len(links),
            'new_links_processed': len(unprocessed_urls),
            'qr_data': qr_data,
            'url_data_csv': csv_path,
            'outputs': {}
        }
        
        if qr_data:
            if self.config.get('output_pptx', True):
                pptx_path = self.output_generator.create_powerpoint(qr_data, file_output_dir)
                if pptx_path:
                    results['outputs']['pptx'] = pptx_path
            
            if self.config.get('output_pdf', True):
                pdf_path = self.output_generator.create_pdf(qr_data, file_output_dir)
                if pdf_path:
                    results['outputs']['pdf'] = pdf_path
        
        return results


def main():
    """Main entry point for the command line interface."""
    parser = argparse.ArgumentParser(
        description="Extract links from files, store in CSV, and generate QR codes."
    )
    parser.add_argument(
        "input_file", 
        help="Path to input file (txt, md, csv, pdf, etc.)"
    )
    parser.add_argument(
        "-o", "--output-dir", 
        default="output",
        help="Base directory for outputs (default: output)"
    )
    parser.add_argument(
        "-c", "--config", 
        help="Path to YAML configuration file"
    )
    
    args = parser.parse_args()
    
    # Check if input file exists
    if not os.path.exists(args.input_file):
        print(f"Error: Input file '{args.input_file}' not found.")
        return 1
    
    # Create app instance
    app = LinkProcessorApp(args.config)
    
    # Process file
    result = app.process_file(args.input_file, args.output_dir)
    
    # Print summary
    print("\nProcessing complete!")
    print(f"Processed file: {result['input_file']}")
    print(f"Found {result['links_found']} unique links")
    print(f"Added {result['new_links']} new links to CSV")
    print(f"Output saved to: {result['output_dir']}")
    
    if result['outputs']['csv']:
        print(f"CSV file created: {result['outputs']['csv']}")
    
    if result['outputs']['pdf']:
        print(f"PDF created: {result['outputs']['pdf']}")
    
    return 0


if __name__ == "__main__":
    main()